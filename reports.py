"""Excel and PDF output."""
import io
import os
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

import calc
from core import CLASS_LABEL, LOGO_PATH, f1, fnum

GREEN = "1E3D36"
RED = "C23B2E"
FONT = "Arial"
thin = Side(style="thin", color="C9CFCB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------
def _bytes(wb):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _body_font(ws, first_row, last_row, ncols):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=cell.font.bold)
            cell.border = BORDER


def _autosize(ws, min_w=8, max_w=42):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color=GREEN)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=10, color="555555")


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Students
# ---------------------------------------------------------------------------
STUDENT_HEADERS = [
    "Admission No",
    "Unique Learner Identifier (ULI)",
    "ASSESSMENT NO.",
    "Full Name",
    "Gender",
    "Grade",
    "Stream",
    "Date of Birth",
    "Guardian",
    "Phone",
]


def students_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(STUDENT_HEADERS)
    _style_header(ws, 1, len(STUDENT_HEADERS))
    widths = [16, 28, 22, 30, 10, 14, 12, 16, 26, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Add sample rows
    ws.append(["1001", "ULI-78912-K", "KPSEA-001", "Amina Otieno", "F", "Grade 7", "East", "2012-03-25", "Grace Otieno", "0712345678"])
    ws.append(["1002", "ULI-78913-L", "KPSEA-002", "Brian Kiprono", "M", "Grade 7", "East", "2012-07-14", "David Kiprono", "0723456789"])
    _body_font(ws, 2, ws.max_row, len(STUDENT_HEADERS))
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Instructions")
    lines = [
        ("How to fill in the Students upload sheet", True),
        ("Enter one student per row on the Students sheet. Keep the column headings exactly as provided.", False),
        ("Admission No and Full Name are required. Admission numbers must be unique across the school.", False),
        ("Unique Learner Identifier (ULI): Government NEMIS / UPI / ULI code for CBC learners.", False),
        ("ASSESSMENT NO.: National assessment / CBA index number (e.g. KPSEA, KCPE index number).", False),
        ("Grade: Class/Grade name (e.g. Grade 7, Form 1, Grade 8).", False),
        ("Stream: Stream name (e.g. East, West, Blue, or leave blank if single stream).", False),
        ("Gender: M (Male) or F (Female).", False),
        ("Date of Birth: Format as YYYY-MM-DD (e.g. 2012-03-25) or standard date.", False),
        ("Guardian: Parent or guardian's full name.", False),
        ("Phone: Contact telephone / mobile number.", False),
        ("", False),
        ("Required Header Layout:", True),
        ("Admission No | Unique Learner Identifier (ULI) | ASSESSMENT NO. | Full Name | Gender | Grade | Stream | Date of Birth | Guardian | Phone", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        ins.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=bold, size=11 if bold else 10)
    ins.column_dimensions["A"].width = 115
    return _bytes(wb)


def students_template_csv():
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(STUDENT_HEADERS)
    writer.writerow(["1001", "ULI-78912-K", "KPSEA-001", "Amina Otieno", "F", "Grade 7", "East", "2012-03-25", "Grace Otieno", "0712345678"])
    writer.writerow(["1002", "ULI-78913-L", "KPSEA-002", "Brian Kiprono", "M", "Grade 7", "East", "2012-07-14", "David Kiprono", "0723456789"])
    return io.BytesIO(buf.getvalue().encode("utf-8-sig"))


def students_workbook(con, class_id=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    headers = STUDENT_HEADERS + ["Status"]
    ws.append(headers)
    sql = ("SELECT s.*, c.name AS cname, c.stream AS cstream FROM students s "
           "LEFT JOIN classes c ON c.id=s.class_id")
    args = ()
    if class_id:
        sql += " WHERE s.class_id=?"
        args = (class_id,)
    sql += " ORDER BY c.name, c.stream, s.full_name"
    for r in con.execute(sql, args):
        ws.append([r["adm_no"], r["uli"] or "", r["assessment_no"] or "", r["full_name"], r["gender"],
                   r["cname"] or "", r["cstream"] or "", r["dob"], r["guardian"], r["phone"], r["status"]])
    _style_header(ws, 1, len(headers))
    _body_font(ws, 2, ws.max_row, len(headers))
    ws.freeze_panes = "A2"
    _autosize(ws)
    return _bytes(wb)


# ---------------------------------------------------------------------------
# Teachers / Users
# ---------------------------------------------------------------------------
TEACHER_HEADERS = ["Full Name", "TSC Number", "Username", "Phone", "Initials", "Role", "Temporary Password"]


def teachers_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Teachers"
    ws.append(TEACHER_HEADERS)
    _style_header(ws, 1, len(TEACHER_HEADERS))
    widths = [28, 16, 18, 18, 12, 18, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # Add sample rows
    ws.append(["Grace Wanjiku", "123456", "gwanjiku", "0712345678", "GW", "Teacher", "Teacher@2026"])
    ws.append(["Peter Kamau", "234567", "pkamau", "0722334455", "PK", "Teacher", "Teacher@2026"])
    ws.append(["David Ochieng", "345678", "dochieng", "0733445566", "DO", "Exam Officer", "Officer@2026"])
    _body_font(ws, 2, ws.max_row, len(TEACHER_HEADERS))
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Instructions")
    lines = [
        ("How to fill in the Teachers upload sheet", True),
        ("Enter one teacher / staff member per row on the Teachers sheet.", False),
        ("Full Name: Teacher's full name (Required, e.g. Grace Wanjiku).", False),
        ("TSC Number: Teachers Service Commission registration number (e.g. 123456).", False),
        ("Username: System login username (Optional. If blank, system automatically creates one like gwanjiku).", False),
        ("Phone: Contact telephone / mobile number (Optional, e.g. 0712345678).", False),
        ("Initials: Staff code or initials shown on report cards (Optional. If blank, auto-generated from name e.g. GW).", False),
        ("Role: System access level (Optional. Options: Teacher, Exam Officer, Viewer, Administrator. Defaults to Teacher).", False),
        ("Temporary Password: Initial password for logging in (Optional. If left blank, the system creates one).", False),
        ("Note: Each teacher will be automatically required to set their own new password upon first login.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        ins.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=bold, size=11 if bold else 10)
    ins.column_dimensions["A"].width = 115
    return _bytes(wb)


def teachers_template_csv():
    import csv
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEACHER_HEADERS)
    writer.writerow(["Grace Wanjiku", "123456", "gwanjiku", "0712345678", "GW", "Teacher", "Teacher@2026"])
    writer.writerow(["Peter Kamau", "234567", "pkamau", "0722334455", "PK", "Teacher", "Teacher@2026"])
    writer.writerow(["David Ochieng", "345678", "dochieng", "0733445566", "DO", "Exam Officer", "Officer@2026"])
    return io.BytesIO(buf.getvalue().encode("utf-8-sig"))


def users_workbook(con):
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    headers = ["Username", "Full Name", "TSC No.", "Initials", "Role", "Phone", "Status", "Teaching Assignments"]
    ws.append(headers)
    rows = con.execute("SELECT u.*, (SELECT COUNT(*) FROM teaching t WHERE t.teacher_id=u.id) AS n_assign "
                       "FROM users u ORDER BY u.active DESC, u.role, u.full_name").fetchall()
    for r in rows:
        role_label = {"teacher": "Teacher", "head_teacher": "Head Teacher / Principal",
                      "deputy_head_teacher": "Deputy Head Teacher", "exam_officer": "Examination Officer",
                      "admin": "Administrator", "viewer": "Viewer"}.get(r["role"], r["role"])
        ws.append([r["username"], r["full_name"], r["tsc_no"] or "", r["initials"] or "", role_label,
                   r["phone"] or "", "Active" if r["active"] else "Disabled", r["n_assign"]])
    _style_header(ws, 1, len(headers))
    _body_font(ws, 2, ws.max_row, len(headers))
    ws.freeze_panes = "A2"
    _autosize(ws)
    return _bytes(wb)


def teachers_credentials_workbook(creds_list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Login Credentials"
    headers = ["#", "Full Name", "TSC No.", "Username", "Temporary Password", "Role", "Phone", "Status"]
    _title(ws, "Teacher Login Credentials", "Important: Teachers must set their own password upon first login.")
    ws.append(headers)
    _style_header(ws, 3, len(headers))
    for i, c in enumerate(creds_list, 1):
        ws.append([i, c.get("name", ""), c.get("tsc_no", ""), c.get("username", ""), c.get("password", ""),
                   c.get("role", ""), c.get("phone", ""), c.get("status", "")])
    _body_font(ws, 4, ws.max_row, len(headers))
    ws.freeze_panes = "A4"
    _autosize(ws)
    return _bytes(wb)


# ---------------------------------------------------------------------------
# Marks
# ---------------------------------------------------------------------------
def marks_template(exam, class_label, students, subj_cols, existing):
    """One sheet per exam and class. subj_cols: [{'code','name','out_of','id'}]."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    headers = ["Adm No", "Student Name"] + [c["code"] for c in subj_cols]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for i, c in enumerate(subj_cols, 3):
        ws.cell(row=1, column=i).comment = Comment(f"{c['name']} - out of {fnum(c['out_of'])}", "System")
    grey = PatternFill("solid", fgColor="EEF1EF")
    for s in students:
        row = [s["adm_no"], s["full_name"]]
        for c in subj_cols:
            score, absent = existing.get((s["id"], c["id"]), (None, 0))
            row.append("ABS" if absent else score)
        ws.append(row)
    _body_font(ws, 2, ws.max_row, len(headers))
    for r in range(2, ws.max_row + 1):
        for cidx in (1, 2):
            ws.cell(row=r, column=cidx).fill = grey
        for cidx in range(3, len(headers) + 1):
            ws.cell(row=r, column=cidx).alignment = Alignment(horizontal="center")
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 34
    for i in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ins = wb.create_sheet("Instructions")
    rows = [("Marks upload sheet", True),
            (f"Examination: {exam['name']} {exam['term']} {exam['year'] or ''}".strip(), False),
            (f"Class: {class_label}", False), ("", False),
            ("Type each student's mark under the subject column. Do not change the Adm No column or the column headings.", False),
            ("Type ABS for a student who was absent. Leave a cell empty to leave that mark unchanged.", False),
            ("Marks may have decimals (for example 45.5) and must not be more than the 'out of' value below.", False),
            ("", False), ("Subject columns in this file", True)]
    for i, (t, b) in enumerate(rows, 1):
        ins.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=b, size=11 if b else 10)
    r0 = len(rows) + 1
    for j, h in enumerate(["Code", "Subject", "Out of"], 1):
        ins.cell(row=r0, column=j, value=h)
    _style_header(ins, r0, 3)
    for i, c in enumerate(subj_cols, r0 + 1):
        ins.cell(row=i, column=1, value=c["code"])
        ins.cell(row=i, column=2, value=c["name"])
        ins.cell(row=i, column=3, value=c["out_of"])
    _body_font(ins, r0 + 1, r0 + len(subj_cols), 3)
    ins.column_dimensions["A"].width = 14
    ins.column_dimensions["B"].width = 40
    ins.column_dimensions["C"].width = 10
    return _bytes(wb)


def marks_workbook(con, exam_id=None):
    """Every recorded mark as a flat table (one row per student per subject)."""
    bands = calc.get_bands(con)
    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    headers = ["Exam", "Term", "Year", "Class", "Adm No", "Student", "Subject code", "Subject",
               "Score", "Out of", "%", "Grade", "Absent", "Status"]
    ws.append(headers)
    sql = (f"SELECT e.name en, e.term, e.year, {CLASS_LABEL} AS clabel, s.adm_no, s.full_name, "
           "sb.code, sb.name sn, m.score, m.absent, m.status, p.out_of "
           "FROM marks m JOIN exams e ON e.id=m.exam_id JOIN students s ON s.id=m.student_id "
           "JOIN classes c ON c.id=m.class_id JOIN subjects sb ON sb.id=m.subject_id "
           "LEFT JOIN exam_papers p ON p.exam_id=m.exam_id AND p.class_id=m.class_id AND p.subject_id=m.subject_id")
    args = ()
    if exam_id:
        sql += " WHERE m.exam_id=?"
        args = (exam_id,)
    sql += " ORDER BY e.year, e.id, c.name, c.stream, s.full_name, sb.name"
    for r in con.execute(sql, args):
        pct = (r["score"] / r["out_of"] * 100) if (r["score"] is not None and r["out_of"]) else None
        b = calc.band_for(pct, bands) if pct is not None else None
        ws.append([r["en"], r["term"], r["year"], r["clabel"], r["adm_no"], r["full_name"], r["code"],
                   r["sn"], r["score"], r["out_of"], round(pct, 1) if pct is not None else None,
                   b["grade"] if b else "", "Yes" if r["absent"] else "", r["status"]])
    _style_header(ws, 1, len(headers))
    _body_font(ws, 2, ws.max_row, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)
    return _bytes(wb)


# ---------------------------------------------------------------------------
# Result tables (Excel)
# ---------------------------------------------------------------------------
def _table_sheet(ws, headers, rows, footer=None, start=4, pct_cols=(), left_cols=()):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start, column=j, value=h)
    _style_header(ws, start, len(headers))
    for i, row in enumerate(rows, start + 1):
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if j - 1 in pct_cols and isinstance(v, (int, float)):
                cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="left" if (j - 1) in left_cols else "center")
    last = start + len(rows)
    if footer:
        last += 1
        for j, v in enumerate(footer, 1):
            cell = ws.cell(row=last, column=j, value=v)
            cell.font = Font(name=FONT, bold=True)
            cell.fill = PatternFill("solid", fgColor="EEF1EF")
            if j - 1 in pct_cols and isinstance(v, (int, float)):
                cell.number_format = "0.0"
            elif isinstance(v, float):
                cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="left" if (j - 1) in left_cols else "center")
    _body_font(ws, start + 1, last, len(headers))
    for r in range(start + 1, last + 1):
        for j in range(1, len(headers) + 1):
            c = ws.cell(row=r, column=j)
            c.font = Font(name=FONT, size=10, bold=c.font.bold)
    ws.freeze_panes = ws.cell(row=start + 1, column=4)
    _autosize(ws)


def merit_xlsx(res, mt, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Class results"
    ex_ = res["exam"]
    _title(ws, title, f"{ex_['name']} {ex_['term']} {ex_['year'] or ''}".strip())
    h = mt["headers"]
    pct = {h.index("Avg %"), h.index("Points")}
    _table_sheet(ws, h, mt["rows"], mt["footer"], pct_cols=pct, left_cols={2, 3} if not mt["single"] else {2})
    return _bytes(wb)


def subjects_xlsx(res, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Subject results"
    ex_ = res["exam"]
    _title(ws, title, f"{ex_['name']} {ex_['term']} {ex_['year'] or ''}".strip())
    grades = [b["grade"] for b in res["bands"]]
    headers = ["Class", "Subject", "Teacher", "Marked", "Absent", "Mean marks", "Mean %", "Highest",
               "Lowest", "Pass rate %"] + grades
    body = []
    for r in rows:
        body.append([r["class"], r["subject"], r["teacher"], r["sat"], r["absent"],
                     r["mean_score"], r["mean_pct"], r["high"], r["low"], r["pass_rate"]]
                    + [r["dist"].get(g, 0) for g in grades])
    _table_sheet(ws, headers, body, pct_cols={5, 6, 9}, left_cols={0, 1, 2})
    return _bytes(wb)


def history_xlsx(student, hist, class_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "History"
    sub = class_label
    extras = []
    if student.get("uli"):
        extras.append(f"ULI: {student['uli']}")
    if student.get("assessment_no"):
        extras.append(f"Assessment No: {student['assessment_no']}")
    if extras:
        sub += "   ·   " + "   ·   ".join(extras)
    _title(ws, f"Performance history: {student['full_name']} ({student['adm_no']})", sub)
    headers = ["Exam", "Term", "Year", "Class", "Total", "Out of", "Avg %", "Grade", "Points",
               "Pos (stream)", "Pos (class)"]
    body = []
    for h in hist:
        s, e = h["student"], h["exam"]
        body.append([e["name"], e["term"], e["year"], h["class_label"], s["total"], s["possible"],
                     s["avg_pct"], s["grade"], s["mean_points"],
                     f"{s['pos_stream']} of {s['n_stream']}", f"{s['pos_class']} of {s['n_class']}"])
    _table_sheet(ws, headers, body, pct_cols={6, 8}, left_cols={0, 1, 3})
    return _bytes(wb)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
INK = colors.HexColor("#1B2321")
CHALK = colors.HexColor("#1E3D36")
PEN = colors.HexColor("#C23B2E")
RULE = colors.HexColor("#C9CFCB")
TINT = colors.HexColor("#F1F4F2")


def _style(name, **kw):
    base = dict(fontName="Helvetica", fontSize=9, leading=11, textColor=INK)
    base.update(kw)
    return ParagraphStyle(name, **base)


def _footer_factory(school, label):
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6B7671"))
        w, h = doc.pagesize
        canvas.drawString(doc.leftMargin, 8 * mm, f"{school} - {label}")
        canvas.drawRightString(w - doc.rightMargin, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()
    return footer


def table_pdf(school, title, lines, headers, rows, footer=None, orient="landscape",
              left_headers=("Name", "Subject", "Teacher", "Class", "Exam", "Term")):
    buf = io.BytesIO()
    page = landscape(A4) if orient == "landscape" else A4
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=title)
    avail = page[0] - 24 * mm
    story = [Paragraph(escape(school), _style("s", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=CHALK)),
             Paragraph(escape(title), _style("t", fontName="Helvetica-Bold", fontSize=11, leading=14))]
    for ln in lines:
        story.append(Paragraph(escape(ln), _style("l", fontSize=8.5, textColor=colors.HexColor("#4A5550"))))
    story.append(Spacer(1, 4 * mm))

    def cell_text(v, h):
        if v is None or v == "":
            return ""
        if isinstance(v, float):
            return f1(v) if h in ("Avg %", "Points", "Mean %", "Pass rate %", "Mean marks") else fnum(v)
        return str(v)

    data = [list(headers)]
    for r in rows:
        data.append([cell_text(v, headers[i]) for i, v in enumerate(r)])
    if footer:
        data.append([cell_text(v, headers[i]) for i, v in enumerate(footer)])
    fs = 8
    while True:
        widths = []
        for j in range(len(headers)):
            w = max(stringWidth(str(row[j]), "Helvetica-Bold" if i == 0 else "Helvetica", fs)
                    for i, row in enumerate(data)) + 7
            widths.append(max(w, 16))
        if sum(widths) <= avail or fs <= 5.5:
            break
        fs -= 0.5
    if sum(widths) > avail:
        widths = [w * avail / sum(widths) for w in widths]
    t = Table(data, colWidths=widths, repeatRows=1)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", fs),
          ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", fs),
          ("BACKGROUND", (0, 0), (-1, 0), CHALK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
          ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("GRID", (0, 0), (-1, -1), 0.4, RULE), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
          ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, TINT])]
    for j, h in enumerate(headers):
        if h in left_headers:
            st.append(("ALIGN", (j, 0), (j, -1), "LEFT"))
    if footer:
        st += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", fs), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3E9E5"))]
    t.setStyle(TableStyle(st))
    story.append(t)
    f = _footer_factory(school, title)
    doc.build(story, onFirstPage=f, onLaterPages=f)
    buf.seek(0)
    return buf


def merit_pdf(school, res, mt, title):
    ex_ = res["exam"]
    lines = [f"{ex_['name']} {ex_['term']} {ex_['year'] or ''}".strip()]
    return table_pdf(school, title, lines, mt["headers"], mt["rows"], mt["footer"])


def subjects_pdf(school, res, rows, title):
    ex_ = res["exam"]
    grades = [b["grade"] for b in res["bands"]]
    headers = ["Class", "Subject", "Teacher", "Marked", "Absent", "Mean marks", "Mean %", "Highest",
               "Lowest", "Pass rate %"] + grades
    body = [[r["class"], r["subject"], r["teacher"], r["sat"], r["absent"], r["mean_score"], r["mean_pct"],
             r["high"], r["low"], r["pass_rate"]] + [r["dist"].get(g, 0) for g in grades] for r in rows]
    return table_pdf(school, title, [f"{ex_['name']} {ex_['term']} {ex_['year'] or ''}".strip()], headers, body)


def history_pdf(school, student, hist, class_label):
    headers = ["Exam", "Term", "Year", "Class", "Total", "Out of", "Avg %", "Grade", "Points",
               "Pos (stream)", "Pos (class)"]
    body = []
    for h in hist:
        s, e = h["student"], h["exam"]
        body.append([e["name"], e["term"], e["year"] or "", h["class_label"], s["total"], s["possible"],
                     s["avg_pct"], s["grade"], s["mean_points"],
                     f"{s['pos_stream']} of {s['n_stream']}", f"{s['pos_class']} of {s['n_class']}"])
    subtitles = [class_label]
    extras = []
    if student.get("uli"):
        extras.append(f"ULI: {student['uli']}")
    if student.get("assessment_no"):
        extras.append(f"Assessment No: {student['assessment_no']}")
    if extras:
        subtitles.append("   ·   ".join(extras))
    return table_pdf(school, f"Performance history: {student['full_name']} ({student['adm_no']})",
                     subtitles, headers, body)


# ---- report cards ---------------------------------------------------------
def _logo_flowable(max_mm=22):
    if not os.path.exists(LOGO_PATH):
        return ""
    try:
        w, h = ImageReader(LOGO_PATH).getSize()
        scale = min(max_mm * mm / w, max_mm * mm / h)
        return Image(LOGO_PATH, width=w * scale, height=h * scale)
    except Exception:
        return ""


def _grade_key(bands):
    return "   ".join(f"{b['grade']} {fnum(b['min_pct'])}+" for b in bands)


def report_cards_pdf(contexts, st):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=11 * mm, bottomMargin=13 * mm, title="Report cards")
    avail = A4[0] - 28 * mm
    story = []
    for i, ctx in enumerate(contexts):
        story += _card(ctx, st, avail)
        if i < len(contexts) - 1:
            story.append(PageBreak())
    f = _footer_factory(st.get("school_name", ""), "Report card")
    doc.build(story, onFirstPage=f, onLaterPages=f)
    buf.seek(0)
    return buf


def _card(ctx, st, avail):
    s, ex_ = ctx["student"], ctx["exam"]
    show_pos = st.get("show_subject_position") == "1"
    P = lambda t, **kw: Paragraph(escape(str(t)), _style("x", **kw))
    out = []

    # header
    school = [P(st.get("school_name", ""), fontName="Helvetica-Bold", fontSize=16, leading=19,
                alignment=1, textColor=CHALK)]
    if st.get("school_motto"):
        school.append(P(st["school_motto"], fontName="Helvetica-Oblique", alignment=1, fontSize=9))
    contact = "   ".join(x for x in [st.get("school_address"), st.get("school_phone"), st.get("school_email")] if x)
    if contact:
        school.append(P(contact, alignment=1, fontSize=8, textColor=colors.HexColor("#4A5550")))
    head = Table([[_logo_flowable(), school, ""]], colWidths=[26 * mm, avail - 52 * mm, 26 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "LEFT")]))
    out += [head, Spacer(1, 2 * mm)]

    title = f"{ex_['name']}" + (f" - {ex_['term']}" if ex_["term"] else "") + (f" {ex_['year']}" if ex_["year"] else "")
    bar = Table([[P("Student report card", fontName="Helvetica-Bold", fontSize=11, textColor=colors.white),
                  P(title, fontName="Helvetica-Bold", fontSize=10, textColor=colors.white, alignment=2)]],
                colWidths=[avail * 0.4, avail * 0.6])
    bar.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), CHALK), ("TOPPADDING", (0, 0), (-1, -1), 5),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 6),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
    out += [bar, Spacer(1, 2.5 * mm)]

    lab = lambda t: P(t, fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.HexColor("#4A5550"))
    val = lambda t: P(t, fontName="Helvetica-Bold", fontSize=10)
    info_rows = [
        [lab("Name"), val(s["name"]), lab("Admission No"), val(s["adm_no"])],
        [lab("Class"), val(ctx["class_label"]), lab("Gender"), val({"M": "Male", "F": "Female"}.get(s["gender"], s["gender"]))]
    ]
    if s.get("uli") or s.get("assessment_no"):
        info_rows.append([
            lab("ULI"), val(s.get("uli") or "-"),
            lab("Assessment No"), val(s.get("assessment_no") or "-")
        ])
    info = Table(info_rows, colWidths=[22 * mm, avail / 2 - 22 * mm, 30 * mm, avail / 2 - 30 * mm])
    info.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                              ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    out += [info, Spacer(1, 3 * mm)]

    # marks table
    hdr = ["Subject", "Marks", "Out of", "%", "Grade"] + (["Pos"] if show_pos else []) + ["Remark", "Teacher"]
    data = [hdr]
    fail_rows = []
    for idx, r in enumerate(ctx["rows"], 1):
        absent = r["absent"]
        row = [r["subject"], "ABS" if absent else fnum(r["score"]), fnum(r["out_of"]),
               "-" if absent else f1(r["pct"]), "-" if absent else r["grade"]]
        if show_pos:
            row.append(r["pos"] if r["pos"] else "-")
        row += ["Absent" if absent else r["remark"], r["teacher"]]
        data.append(row)
        if r["pct"] is not None and r["pct"] < ctx["pass_mark"]:
            fail_rows.append(idx)
    w = [50, 16, 16, 16, 15] + ([13] if show_pos else []) + [34 if show_pos else 47, 22]
    scale = avail / (sum(w) * mm)
    widths = [x * mm * scale for x in w]
    t = Table(data, colWidths=widths, repeatRows=1)
    ts = [("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
          ("BACKGROUND", (0, 0), (-1, 0), TINT), ("LINEBELOW", (0, 0), (-1, 0), 1, CHALK),
          ("LINEBELOW", (0, 1), (-1, -1), 0.4, RULE), ("ALIGN", (1, 0), (-1, -1), "CENTER"),
          ("ALIGN", (-2, 0), (-1, -1), "LEFT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("TOPPADDING", (0, 0), (-1, -1), 3.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.2),
          ("FONT", (4, 1), (4, -1), "Helvetica-Bold", 9.5), ("TEXTCOLOR", (0, 0), (-1, -1), INK)]
    for idx in fail_rows:
        ts.append(("TEXTCOLOR", (3, idx), (4, idx), PEN))
    t.setStyle(TableStyle(ts))
    out += [t, Spacer(1, 3 * mm)]

    # summary
    def cellp(label, value):
        return [P(label, fontSize=7.5, textColor=colors.HexColor("#4A5550")),
                P(value, fontName="Helvetica-Bold", fontSize=12, leading=14)]
    pos_s = f"{s['pos_stream']} of {s['n_stream']}" if s["pos_stream"] else "-"
    pos_c = f"{s['pos_class']} of {s['n_class']}" if s["pos_class"] else "-"
    cells = [cellp("Total marks", f"{fnum(s['total'])} / {fnum(s['possible'])}"),
             cellp("Average", f"{f1(s['avg_pct'])}%" if s["n"] else "-"),
             cellp("Mean grade", s["grade"] or "-"),
             cellp("Mean points", f1(s["mean_points"]) if s["n"] else "-"),
             cellp("Position in stream", pos_s), cellp("Position in class", pos_c)]
    summ = Table([cells], colWidths=[avail / 6] * 6)
    summ.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, CHALK), ("INNERGRID", (0, 0), (-1, -1), 0.4, RULE),
                              ("BACKGROUND", (0, 0), (-1, -1), TINT), ("TOPPADDING", (0, 0), (-1, -1), 4),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    out += [summ, Spacer(1, 1.5 * mm),
            P("Grade key (percentage):  " + _grade_key(ctx["bands"]), fontSize=7, textColor=colors.HexColor("#4A5550")),
            Spacer(1, 3 * mm)]

    # remarks and signatures
    def remark_box(label, text, who=""):
        body = [P(label, fontName="Helvetica-Bold", fontSize=8.5),
                P(text or " ", fontSize=9.5, leading=12),
                Spacer(1, 3 * mm),
                P("Signature: ______________________     Date: ______________" + (f"     ({who})" if who else ""),
                  fontSize=8, textColor=colors.HexColor("#4A5550"))]
        tb = Table([[body]], colWidths=[avail], rowHeights=[22 * mm])
        tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 4)]))
        return tb
    out.append(remark_box("Class teacher's remarks", ctx["teacher_remark"]))
    out.append(Spacer(1, 2 * mm))
    out.append(remark_box("Principal's remarks", ctx["principal_remark"], st.get("principal_name", "")))
    dates = []
    if ex_.get("closing_date"):
        dates.append(f"Closing date: {ex_['closing_date']}")
    if ex_.get("opening_date"):
        dates.append(f"Next term opens: {ex_['opening_date']}")
    if dates:
        out += [Spacer(1, 2.5 * mm), P("     ".join(dates), fontName="Helvetica-Bold", fontSize=9)]
    if st.get("report_footer"):
        out += [Spacer(1, 2 * mm), P(st["report_footer"], fontName="Helvetica-Oblique", fontSize=7.5,
                                     alignment=1, textColor=colors.HexColor("#4A5550"))]
    return out
