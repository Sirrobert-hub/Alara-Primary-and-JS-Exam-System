"""Examinations, marks entry, Excel marks upload, and verification."""
import datetime
import io
import os
import re
import secrets

import openpyxl
from flask import Blueprint, abort, flash, g, redirect, render_template, request, send_file, url_for

import calc
import reports
from core import (CLASS_LABEL, TMP_DIR, audit, can_enter_marks, commit, db, ex, flash_errors, fnum,
                  has_perm, is_admin, login_required, norm_cell, now, perm_required, q, q1, teacher_papers,
                  to_float)

bp = Blueprint("exams", __name__)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ABSENT_WORDS = {"ABS", "A", "AB", "ABSENT"}


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def get_exam(eid):
    e = q1("SELECT * FROM exams WHERE id=?", (eid,))
    if not e:
        abort(404)
    return e


def paper_progress(eid):
    return q(f"""
        SELECT p.*, {CLASS_LABEL} AS clabel, sb.code, sb.name AS sname, u.full_name AS teacher, u.tsc_no AS teacher_tsc,
          (SELECT COUNT(*) FROM students s WHERE s.class_id=p.class_id AND s.status='active') AS expected,
          (SELECT COUNT(*) FROM marks m WHERE m.exam_id=p.exam_id AND m.class_id=p.class_id
              AND m.subject_id=p.subject_id AND (m.score IS NOT NULL OR m.absent=1)) AS entered,
          (SELECT COUNT(*) FROM marks m WHERE m.exam_id=p.exam_id AND m.class_id=p.class_id
              AND m.subject_id=p.subject_id AND m.status='verified') AS verified
        FROM exam_papers p
        JOIN classes c ON c.id=p.class_id JOIN subjects sb ON sb.id=p.subject_id
        LEFT JOIN teaching t ON t.class_id=p.class_id AND t.subject_id=p.subject_id
        LEFT JOIN users u ON u.id=t.teacher_id
        WHERE p.exam_id=? ORDER BY c.name, c.stream, sb.name""", (eid,))


def can_override():
    return is_admin() or has_perm("marks.verify")


def mark_locked(exam, mark_status):
    """Teachers cannot change marks once the exam is closed or the mark is verified."""
    if can_override():
        return False
    return exam["status"] != "open" or mark_status == "verified"


def apply_marks(exam, cid, subject_id, changes, adm_by_id):
    """changes: [(student_id, score|None, absent bool)]. score None and not absent => remove."""
    if not can_enter_marks(cid, subject_id):
        abort(403)
    # Ensure students belong to this class and are active (only students taught in this class)
    valid_student_ids = {r["id"] for r in q(
        "SELECT id FROM students WHERE class_id=? AND status='active'", (cid,))}
    filtered_changes = [ch for ch in changes if ch[0] in valid_student_ids]
    if g.user and g.user["role"] == "teacher" and len(filtered_changes) != len(changes):
        # Teacher attempted to submit marks for students outside their assigned class
        abort(403)
    changes = filtered_changes

    sub = q1("SELECT code FROM subjects WHERE id=?", (subject_id,))["code"]
    existing = {r["student_id"]: r for r in q(
        "SELECT * FROM marks WHERE exam_id=? AND subject_id=?", (exam["id"], subject_id))}
    added = updated = removed = reverify = 0
    lines = []
    for stid, score, absent in changes:
        old = existing.get(stid)
        if mark_locked(exam, old["status"] if old else None):
            continue
        adm = adm_by_id.get(stid, str(stid))
        if score is None and not absent:
            if old:
                ex("DELETE FROM marks WHERE id=?", (old["id"],))
                removed += 1
                lines.append(f"{adm} {sub}: removed (was {'ABS' if old['absent'] else fnum(old['score'])})")
            continue
        if old is None:
            ex("INSERT INTO marks(exam_id,student_id,class_id,subject_id,score,absent,status,entered_by,updated_at) "
               "VALUES(?,?,?,?,?,?,'entered',?,?)",
               (exam["id"], stid, cid, subject_id, score, 1 if absent else 0, g.user["id"], now()))
            added += 1
        else:
            if old["score"] == score and bool(old["absent"]) == absent and old["class_id"] == cid:
                continue
            ex("UPDATE marks SET class_id=?, score=?, absent=?, status='entered', entered_by=?, updated_at=? WHERE id=?",
               (cid, score, 1 if absent else 0, g.user["id"], now(), old["id"]))
            updated += 1
            if old["status"] == "verified":
                reverify += 1
            was = "ABS" if old["absent"] else fnum(old["score"])
            lines.append(f"{adm} {sub}: {was} -> {'ABS' if absent else fnum(score)}")
    for ln in lines[:60]:
        audit("mark changed", f"{exam['name']}: {ln}")
    if added or updated or removed:
        audit("marks saved", f"{exam['name']} / {sub}: {added} added, {updated} changed, {removed} removed")
    return added, updated, removed, reverify


# ---------------------------------------------------------------------------
# examinations
# ---------------------------------------------------------------------------
@bp.route("/exams")
@login_required
def exams():
    rows = []
    for e in q("SELECT * FROM exams ORDER BY year DESC, id DESC"):
        n = q1("SELECT COUNT(*) papers, COUNT(DISTINCT class_id) classes FROM exam_papers WHERE exam_id=?", (e["id"],))
        exp = q1("SELECT COALESCE(SUM((SELECT COUNT(*) FROM students s WHERE s.class_id=p.class_id AND s.status='active')),0) n "
                 "FROM exam_papers p WHERE p.exam_id=?", (e["id"],))["n"]
        got = q1("SELECT COUNT(*) n FROM marks WHERE exam_id=? AND (score IS NOT NULL OR absent=1)", (e["id"],))["n"]
        rows.append({**dict(e), "papers": n["papers"], "classes": n["classes"],
                     "pct": min(round(got / exp * 100), 100) if exp else 0})
    return render_template("exams.html", exams=rows)


def _exam_form_data():
    f = request.form
    return dict(name=f.get("name", "").strip(), term=f.get("term", "").strip(),
                year=f.get("year", type=int), out_of=to_float(f.get("out_of")) or 100.0,
                closing_date=f.get("closing_date", "").strip(), opening_date=f.get("opening_date", "").strip())


@bp.route("/exams/new", methods=["GET", "POST"])
@bp.route("/exams/<int:eid>/edit", methods=["GET", "POST"])
@perm_required("exams.manage")
def exam_form(eid=None):
    exam = get_exam(eid) if eid else None
    classes = q(f"SELECT c.id, {CLASS_LABEL} AS label FROM classes c ORDER BY c.name, c.stream")
    subjects = q("SELECT * FROM subjects ORDER BY name")
    if request.method == "POST":
        d = _exam_form_data()
        cls_ids = request.form.getlist("classes", type=int)
        sub_ids = request.form.getlist("subjects", type=int)
        errors = []
        if not d["name"]:
            errors.append("Give the examination a name, for example End of Term 1 Exam.")
        if d["out_of"] <= 0:
            errors.append("Marks 'out of' must be more than zero.")
        if not exam and (not cls_ids or not sub_ids):
            errors.append("Choose at least one class and one subject.")
        if errors:
            flash_errors(errors)
            return render_template("exam_form.html", exam=d, eid=eid, classes=classes, subjects=subjects,
                                   sel_c=set(cls_ids), sel_s=set(sub_ids))
        if exam:
            ex("UPDATE exams SET name=?, term=?, year=?, out_of=?, closing_date=?, opening_date=? WHERE id=?",
               (d["name"], d["term"], d["year"], d["out_of"], d["closing_date"], d["opening_date"], eid))
            audit("exam edited", d["name"])
            commit()
            flash("Examination updated.", "ok")
            return redirect(url_for("exams.exam_detail", eid=eid))
        cur = ex("INSERT INTO exams(name,term,year,out_of,closing_date,opening_date) VALUES(?,?,?,?,?,?)",
                 (d["name"], d["term"], d["year"], d["out_of"], d["closing_date"], d["opening_date"]))
        new_id = cur.lastrowid
        for c in cls_ids:
            for s in sub_ids:
                ex("INSERT INTO exam_papers(exam_id,class_id,subject_id,out_of) VALUES(?,?,?,?)",
                   (new_id, c, s, d["out_of"]))
        audit("exam created", f"{d['name']}: {len(cls_ids)} classes, {len(sub_ids)} subjects")
        commit()
        flash("Examination created. You can fine-tune which subjects each class sits on the next screen.", "ok")
        return redirect(url_for("exams.exam_setup", eid=new_id))
    if exam:
        default = dict(exam)
    else:
        default = {"name": "", "term": "", "year": datetime.date.today().year, "out_of": 100,
                   "closing_date": "", "opening_date": ""}
    return render_template("exam_form.html", exam=default, eid=eid, classes=classes, subjects=subjects,
                           sel_c=set(), sel_s=set())


@bp.route("/exams/<int:eid>")
@login_required
def exam_detail(eid):
    exam = get_exam(eid)
    papers = paper_progress(eid)
    see_all = has_perm("exams.manage") or has_perm("marks.enter_all") or has_perm("marks.verify") \
        or has_perm("results.view_all")
    if not see_all:
        mine = teacher_papers()
        papers = [p for p in papers if (p["class_id"], p["subject_id"]) in mine]
    editable = {(p["class_id"], p["subject_id"]) for p in papers if can_enter_marks(p["class_id"], p["subject_id"])}
    return render_template("exam_detail.html", exam=exam, papers=papers, editable=editable, see_all=see_all)


@bp.route("/exams/<int:eid>/setup", methods=["GET", "POST"])
@perm_required("exams.manage")
def exam_setup(eid):
    exam = get_exam(eid)
    classes = q(f"SELECT c.id, {CLASS_LABEL} AS label FROM classes c ORDER BY c.name, c.stream")
    subjects = q("SELECT * FROM subjects ORDER BY name")
    have = {(r["class_id"], r["subject_id"]): r for r in q("SELECT * FROM exam_papers WHERE exam_id=?", (eid,))}
    if request.method == "POST":
        errors = []
        outs = {}
        for s in subjects:
            v = to_float(request.form.get(f"out_{s['id']}")) if request.form.get(f"out_{s['id']}", "").strip() else exam["out_of"]
            if v is None or v <= 0:
                errors.append(f"'Out of' for {s['name']} must be a number above zero.")
                continue
            mx = q1("SELECT MAX(score) m FROM marks WHERE exam_id=? AND subject_id=?", (eid, s["id"]))["m"]
            if mx is not None and mx > v:
                errors.append(f"{s['name']}: a mark of {fnum(mx)} has already been entered, so 'out of' cannot be {fnum(v)}.")
                continue
            outs[s["id"]] = v
        removals = []
        for c in classes:
            for s in subjects:
                want = bool(request.form.get(f"p_{c['id']}_{s['id']}"))
                key = (c["id"], s["id"])
                if want and key not in have and s["id"] in outs:
                    ex("INSERT INTO exam_papers(exam_id,class_id,subject_id,out_of) VALUES(?,?,?,?)",
                       (eid, c["id"], s["id"], outs[s["id"]]))
                elif not want and key in have:
                    if q1("SELECT 1 FROM marks WHERE exam_id=? AND class_id=? AND subject_id=? LIMIT 1",
                          (eid, c["id"], s["id"])) and not is_admin():
                        errors.append(f"{c['label']} / {s['name']} already has marks, so only an administrator can remove it.")
                    else:
                        removals.append(key)
        if errors:
            db().rollback()
            flash_errors(errors)
        else:
            for c_id, s_id in removals:
                ex("DELETE FROM marks WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, c_id, s_id))
                ex("DELETE FROM exam_papers WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, c_id, s_id))
            for s_id, v in outs.items():
                ex("UPDATE exam_papers SET out_of=? WHERE exam_id=? AND subject_id=?", (v, eid, s_id))
            audit("exam setup saved", exam["name"])
            commit()
            flash("Classes and subjects saved.", "ok")
            return redirect(url_for("exams.exam_detail", eid=eid))
        have = {(r["class_id"], r["subject_id"]): r for r in q("SELECT * FROM exam_papers WHERE exam_id=?", (eid,))}
    out_of = {}
    for (c_id, s_id), p in have.items():
        out_of.setdefault(s_id, p["out_of"])
    return render_template("exam_setup.html", exam=exam, classes=classes, subjects=subjects,
                           have=set(have), out_of=out_of)


@bp.route("/exams/<int:eid>/paper/<int:cid>/<int:sid>/delete", methods=["POST"])
@perm_required("exams.manage")
def exam_paper_delete(eid, cid, sid):
    exam = get_exam(eid)
    n = ex("DELETE FROM marks WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, cid, sid)).rowcount
    ex("DELETE FROM exam_papers WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, cid, sid))
    audit("exam paper entry deleted", f"{exam['name']}: class {cid}, subject {sid} ({n} marks deleted)")
    commit()
    flash(f"Exam paper entry deleted and {n} marks removed.", "ok")
    return redirect(url_for("exams.exam_detail", eid=eid))


@bp.route("/exams/<int:eid>/paper/<int:cid>/<int:sid>/clear", methods=["POST"])
@perm_required("exams.manage")
def exam_paper_clear(eid, cid, sid):
    exam = get_exam(eid)
    n = ex("DELETE FROM marks WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, cid, sid)).rowcount
    audit("marks cleared", f"{exam['name']}: class {cid}, subject {sid} ({n} marks cleared)")
    commit()
    flash(f"All {n} marks entries for this paper have been cleared.", "ok")
    return redirect(url_for("exams.marks_entry", eid=eid, cid=cid, sid=sid))


@bp.route("/exams/<int:eid>/status", methods=["POST"])
@perm_required("exams.manage")
def exam_status(eid):
    exam = get_exam(eid)
    new = "closed" if exam["status"] == "open" else "open"
    ex("UPDATE exams SET status=? WHERE id=?", (new, eid))
    audit("exam " + ("closed" if new == "closed" else "reopened"), exam["name"])
    commit()
    flash("Marks entry is now closed for teachers." if new == "closed" else "Marks entry reopened.", "ok")
    return redirect(request.referrer or url_for("exams.exam_detail", eid=eid))


@bp.route("/exams/<int:eid>/delete", methods=["POST"])
@perm_required("exams.manage")
def exam_delete(eid):
    exam = get_exam(eid)
    has_marks = q1("SELECT COUNT(*) n FROM marks WHERE exam_id=?", (eid,))["n"]
    if has_marks and request.form.get("confirm", "").strip().upper() != "DELETE":
        flash(f"This examination has {has_marks} marks. Type DELETE in the box to confirm.", "error")
        return redirect(url_for("exams.exam_detail", eid=eid))
    ex("DELETE FROM marks WHERE exam_id=?", (eid,))
    ex("DELETE FROM report_remarks WHERE exam_id=?", (eid,))
    ex("DELETE FROM exam_papers WHERE exam_id=?", (eid,))
    ex("DELETE FROM exams WHERE id=?", (eid,))
    audit("exam deleted", f"{exam['name']} ({has_marks} marks)")
    commit()
    flash("Examination deleted.", "ok")
    return redirect(url_for("exams.exams"))


@bp.route("/exams/<int:eid>/export")
@perm_required("exams.manage", "results.view_all")
def exam_export(eid):
    exam = get_exam(eid)
    data = reports.marks_workbook(db(), eid)
    return send_file(data, as_attachment=True, download_name=f"marks_{re.sub(r'[^A-Za-z0-9]+', '_', exam['name'])}.xlsx",
                     mimetype=XLSX)


# ---------------------------------------------------------------------------
# manual marks entry
# ---------------------------------------------------------------------------
@bp.route("/marks")
@login_required
def marks_home():
    all_exams = q("SELECT * FROM exams ORDER BY year DESC, id DESC")
    eid = request.args.get("exam_id", type=int)
    if not eid:
        pick = [e for e in all_exams if e["status"] == "open"] or all_exams
        eid = pick[0]["id"] if pick else None
    papers = []
    exam = None
    if eid:
        exam = get_exam(eid)
        papers = [p for p in paper_progress(eid) if can_enter_marks(p["class_id"], p["subject_id"])]
    return render_template("marks_home.html", exams=all_exams, exam=exam, papers=papers)


@bp.route("/marks/enter/<int:eid>/<int:cid>/<int:sid>", methods=["GET", "POST"])
@login_required
def marks_entry(eid, cid, sid):
    exam = get_exam(eid)
    paper = q1("SELECT * FROM exam_papers WHERE exam_id=? AND class_id=? AND subject_id=?", (eid, cid, sid))
    if not paper:
        abort(404)
    if not can_enter_marks(cid, sid):
        abort(403)
    cls = q1(f"SELECT c.id, {CLASS_LABEL} AS label FROM classes c WHERE c.id=?", (cid,))
    subject = q1("SELECT * FROM subjects WHERE id=?", (sid,))
    students = q("SELECT * FROM students WHERE class_id=? AND status='active' ORDER BY full_name", (cid,))
    marks = {m["student_id"]: m for m in q("SELECT * FROM marks WHERE exam_id=? AND subject_id=?", (eid, sid))}
    entered = {}
    if request.method == "POST":
        out_of, changes, errors = paper["out_of"], [], []
        for s in students:
            m = marks.get(s["id"])
            if m and mark_locked(exam, m["status"]):
                continue
            if not m and mark_locked(exam, None):
                continue
            raw = request.form.get(f"s_{s['id']}", "").strip()
            absent = bool(request.form.get(f"a_{s['id']}"))
            entered[s["id"]] = (raw, absent)
            if absent:
                changes.append((s["id"], None, True))
            elif raw == "":
                changes.append((s["id"], None, False))
            else:
                try:
                    val = to_float(raw)
                except ValueError:
                    errors.append(f"{s['full_name']}: \"{raw}\" is not a number.")
                    continue
                if val < 0 or val > out_of:
                    errors.append(f"{s['full_name']}: {raw} is outside 0 to {fnum(out_of)}.")
                    continue
                changes.append((s["id"], round(val, 2), False))
        if errors:
            flash_errors(errors)
            flash("Nothing was saved. Correct the highlighted problems and save again.", "error")
        else:
            a, u, r, rv = apply_marks(exam, cid, sid, changes, {s["id"]: s["adm_no"] for s in students})
            commit()
            msg = f"Saved: {a} new, {u} changed, {r} removed."
            if rv:
                msg += f" {rv} verified marks were changed and need to be verified again."
            flash(msg, "ok")
            return redirect(url_for("exams.marks_entry", eid=eid, cid=cid, sid=sid))
    rows = []
    for s in students:
        m = marks.get(s["id"])
        raw, absent = entered.get(s["id"], (None, None))
        if raw is None:
            raw = fnum(m["score"]) if m and m["score"] is not None else ""
            absent = bool(m and m["absent"])
        rows.append({"s": s, "value": raw, "absent": absent,
                     "locked": mark_locked(exam, m["status"] if m else None),
                     "verified": bool(m and m["status"] == "verified")})
    teacher = q1("SELECT u.full_name, u.tsc_no FROM teaching t JOIN users u ON u.id=t.teacher_id "
                 "WHERE t.class_id=? AND t.subject_id=?", (cid, sid))
    t_label = ""
    if teacher:
        t_label = teacher["full_name"]
        if teacher["tsc_no"]:
            t_label += f" (TSC: {teacher['tsc_no']})"
    return render_template("marks_entry.html", exam=exam, paper=paper, cls=cls, subject=subject, rows=rows,
                           teacher=t_label,
                           all_locked=all(r["locked"] for r in rows) if rows else True,
                           override=can_override())


# ---------------------------------------------------------------------------
# Excel marks upload
# ---------------------------------------------------------------------------
def _permitted_papers(eid, cid):
    rows = q("SELECT p.*, sb.code, sb.name AS sname FROM exam_papers p JOIN subjects sb ON sb.id=p.subject_id "
             "WHERE p.exam_id=? AND p.class_id=? ORDER BY sb.name", (eid, cid))
    return rows, [p for p in rows if can_enter_marks(cid, p["subject_id"])]


@bp.route("/marks/template/<int:eid>/<int:cid>")
@login_required
def marks_template(eid, cid):
    exam = get_exam(eid)
    _all, mine = _permitted_papers(eid, cid)
    if not mine:
        abort(403)
    cls = q1(f"SELECT {CLASS_LABEL} AS label FROM classes c WHERE c.id=?", (cid,))["label"]
    students = q("SELECT * FROM students WHERE class_id=? AND status='active' ORDER BY full_name", (cid,))
    existing = {(m["student_id"], m["subject_id"]): (m["score"], m["absent"]) for m in q(
        "SELECT * FROM marks WHERE exam_id=? AND class_id=?", (eid, cid))}
    cols = [{"id": p["subject_id"], "code": p["code"], "name": p["sname"], "out_of": p["out_of"]} for p in mine]
    data = reports.marks_template(exam, cls, students, cols, existing)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", f"{exam['name']}_{cls}")
    return send_file(data, as_attachment=True, download_name=f"marks_template_{safe}.xlsx", mimetype=XLSX)


def parse_marks_file(raw_bytes, exam, cid):
    """Read an uploaded workbook. Returns dict with changes, errors, warnings and preview rows."""
    result = {"changes": [], "errors": [], "warnings": [], "preview": [], "matched": 0, "skipped_locked": 0,
              "unchanged": 0, "new": 0, "changed": 0}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception:
        result["errors"].append("That file could not be opened. Save it as a normal .xlsx workbook and try again.")
        return result
    ws = wb["Marks"] if "Marks" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    all_papers, mine = _permitted_papers(exam["id"], cid)
    allowed = {p["code"].upper(): p for p in mine}
    every = {p["code"].upper() for p in all_papers}
    header_idx, adm_col = None, None
    for i, row in enumerate(rows[:10]):
        for j, cell in enumerate(row):
            if norm_cell(cell).lower() in ("adm no", "admission no", "admission number", "adm", "admission"):
                header_idx, adm_col = i, j
                break
        if header_idx is not None:
            break
    if header_idx is None:
        result["errors"].append("The 'Adm No' column was not found. Download the template from this page and use it.")
        return result
    sub_cols, ignored, unknown = {}, [], []
    for j, cell in enumerate(rows[header_idx]):
        code = norm_cell(cell).upper()
        if j == adm_col or not code or code in ("STUDENT NAME", "NAME", "FULL NAME"):
            continue
        if code in allowed:
            sub_cols[j] = allowed[code]
        elif code in every:
            ignored.append(code)
        else:
            unknown.append(norm_cell(cell))
    if ignored:
        result["warnings"].append("Columns skipped because you are not assigned to those subjects: " + ", ".join(ignored) + ".")
    if unknown:
        result["warnings"].append("Columns not recognised and ignored: " + ", ".join(unknown) + ".")
    if not sub_cols:
        result["errors"].append("No subject columns you can enter marks for were found in this file.")
        return result
    students = {s["adm_no"].lower(): s for s in q("SELECT * FROM students WHERE class_id=? AND status='active'", (cid,))}
    marks = {(m["student_id"], m["subject_id"]): m for m in q(
        "SELECT * FROM marks WHERE exam_id=? AND class_id=?", (exam["id"], cid))}
    seen = set()
    for n, row in enumerate(rows[header_idx + 1:], header_idx + 2):
        adm = norm_cell(row[adm_col]) if adm_col < len(row) else ""
        if not adm:
            continue
        st = students.get(adm.lower())
        if not st:
            result["errors"].append(f"Row {n}: {adm} is not an active student in this class.")
            continue
        if adm.lower() in seen:
            result["errors"].append(f"Row {n}: {adm} appears more than once.")
            continue
        seen.add(adm.lower())
        result["matched"] += 1
        for j, p in sub_cols.items():
            v = row[j] if j < len(row) else None
            if v is None or str(v).strip() == "":
                continue
            txt = str(v).strip()
            if txt.upper() in ABSENT_WORDS:
                score, absent = None, True
            else:
                try:
                    score, absent = to_float(txt), False
                except ValueError:
                    result["errors"].append(f"Row {n} ({adm}), {p['code']}: \"{txt}\" is not a number.")
                    continue
                if score < 0 or score > p["out_of"]:
                    result["errors"].append(f"Row {n} ({adm}), {p['code']}: {txt} is outside 0 to {fnum(p['out_of'])}.")
                    continue
                score = round(score, 2)
            old = marks.get((st["id"], p["subject_id"]))
            if old and mark_locked(exam, old["status"]):
                result["skipped_locked"] += 1
                continue
            if not old and mark_locked(exam, None):
                result["skipped_locked"] += 1
                continue
            if old and old["score"] == score and bool(old["absent"]) == absent:
                result["unchanged"] += 1
                continue
            result["changes"].append((st["id"], p["subject_id"], score, absent))
            if old:
                result["changed"] += 1
                was = "ABS" if old["absent"] else fnum(old["score"])
            else:
                result["new"] += 1
                was = ""
            if len(result["preview"]) < 40:
                result["preview"].append((st["adm_no"], st["full_name"], p["code"], was, "ABS" if absent else fnum(score)))
    if result["skipped_locked"]:
        result["warnings"].append(f"{result['skipped_locked']} marks were skipped because they are verified or the examination is closed.")
    return result


@bp.route("/marks/import", methods=["GET", "POST"])
@login_required
def marks_import():
    exams_ = q("SELECT * FROM exams ORDER BY year DESC, id DESC")
    eid = request.values.get("exam_id", type=int)
    cid = request.values.get("class_id", type=int)
    classes = []
    if eid:
        for c in q(f"SELECT DISTINCT c.id, {CLASS_LABEL} AS label FROM exam_papers p JOIN classes c ON c.id=p.class_id "
                   "WHERE p.exam_id=? ORDER BY c.name, c.stream", (eid,)):
            if _permitted_papers(eid, c["id"])[1]:
                classes.append(c)
    preview = None
    if request.method == "POST" and eid and cid:
        exam = get_exam(eid)
        if not _permitted_papers(eid, cid)[1]:
            abort(403)
        file = request.files.get("file")
        if not file or not file.filename.lower().endswith((".xlsx", ".xlsm")):
            flash("Choose an Excel file (.xlsx) to upload.", "error")
        else:
            raw = file.read()
            preview = parse_marks_file(raw, exam, cid)
            preview["token"] = None
            if not preview["errors"] and preview["changes"]:
                token = secrets.token_hex(8)
                os.makedirs(TMP_DIR, exist_ok=True)
                with open(os.path.join(TMP_DIR, f"marks_{g.user['id']}_{token}.xlsx"), "wb") as f:
                    f.write(raw)
                preview["token"] = token
    return render_template("marks_import.html", exams=exams_, classes=classes, eid=eid, cid=cid, preview=preview)


@bp.route("/marks/import/confirm", methods=["POST"])
@login_required
def marks_import_confirm():
    eid, cid = request.form.get("exam_id", type=int), request.form.get("class_id", type=int)
    if not eid or not cid:
        abort(400)
    permitted = _permitted_papers(eid, cid)[1]
    if not permitted:
        abort(403)
    token = request.form.get("token", "")
    if not re.fullmatch(r"[0-9a-f]{16}", token):
        abort(400)
    path = os.path.join(TMP_DIR, f"marks_{g.user['id']}_{token}.xlsx")
    if not os.path.exists(path):
        flash("The upload has expired. Please upload the file again.", "error")
        return redirect(url_for("exams.marks_import", exam_id=eid, class_id=cid))
    exam = get_exam(eid)
    with open(path, "rb") as f:
        result = parse_marks_file(f.read(), exam, cid)
    os.remove(path)
    if result["errors"]:
        flash_errors(result["errors"])
        return redirect(url_for("exams.marks_import", exam_id=eid, class_id=cid))
    adm = {s["id"]: s["adm_no"] for s in q("SELECT id, adm_no FROM students WHERE class_id=?", (cid,))}
    by_subject = {}
    for stid, subj, score, absent in result["changes"]:
        by_subject.setdefault(subj, []).append((stid, score, absent))
    a = u = rv = 0
    for subj, ch in by_subject.items():
        if not can_enter_marks(cid, subj):
            continue
        x, y, _, z = apply_marks(exam, cid, subj, ch, adm)
        a, u, rv = a + x, u + y, rv + z
    audit("marks imported", f"{exam['name']}: Excel upload, {a} new, {u} changed")
    commit()
    msg = f"Marks imported: {a} new and {u} changed."
    if rv:
        msg += f" {rv} verified marks were changed and need to be verified again."
    flash(msg, "ok")
    return redirect(url_for("exams.exam_detail", eid=eid))


# ---------------------------------------------------------------------------
# verification
# ---------------------------------------------------------------------------
@bp.route("/verify")
@perm_required("marks.verify")
def verify_home():
    all_exams = q("SELECT * FROM exams ORDER BY year DESC, id DESC")
    eid = request.args.get("exam_id", type=int) or (all_exams[0]["id"] if all_exams else None)
    papers, stats, exam = [], {}, None
    if eid:
        exam = get_exam(eid)
        papers = paper_progress(eid)
        res = calc.compute_results(db(), eid, only_verified=False)
        stats = res["paper_stats"]
    return render_template("verify.html", exams=all_exams, exam=exam, papers=papers, stats=stats)


@bp.route("/verify/<int:eid>/set", methods=["POST"])
@perm_required("marks.verify")
def verify_set(eid):
    exam = get_exam(eid)
    action = request.form.get("action")
    cid, sid = request.form.get("class_id", type=int), request.form.get("subject_id", type=int)
    if action == "verify_all":
        n = 0
        for p in paper_progress(eid):
            if p["expected"] and p["entered"] >= p["expected"] and p["verified"] < p["entered"]:
                n += ex("UPDATE marks SET status='verified' WHERE exam_id=? AND class_id=? AND subject_id=? "
                        "AND status='entered'", (eid, p["class_id"], p["subject_id"])).rowcount
        audit("marks verified", f"{exam['name']}: all complete papers ({n} marks)")
        flash(f"{n} marks verified in all papers that are fully entered.", "ok")
    elif action in ("verify", "unlock") and cid and sid:
        new, old = ("verified", "entered") if action == "verify" else ("entered", "verified")
        n = ex("UPDATE marks SET status=? WHERE exam_id=? AND class_id=? AND subject_id=? AND status=?",
               (new, eid, cid, sid, old)).rowcount
        audit("marks " + ("verified" if action == "verify" else "unlocked"),
              f"{exam['name']}: class {cid}, subject {sid} ({n} marks)")
        flash(f"{n} marks {'verified' if action == 'verify' else 'unlocked for editing'}.", "ok")
    commit()
    return redirect(url_for("exams.verify_home", exam_id=eid))
