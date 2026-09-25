"""Shared helpers: database, sessions, permissions, settings, audit log, backups.

Everything is stored in the ./data folder next to the program:
    data/exams.db      - the database (a single SQLite file)
    data/backups/      - backup archives
    data/secret.key    - key used to sign login sessions
"""
import datetime
import functools
import hmac
import os
import secrets
import shutil
import sqlite3
import zipfile

from flask import abort, flash, g, redirect, request, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("EXAM_DATA_DIR") or os.path.join(BASE_DIR, "data")
DB_PATH = os.path.join(DATA_DIR, "exams.db")
BACKUP_DIR = os.path.join(DATA_DIR, "backups")
TMP_DIR = os.path.join(DATA_DIR, "tmp")
LOGO_PATH = os.path.join(DATA_DIR, "logo.png")

APP_VERSION = "1.0"

# --------------------------------------------------------------------------
# Roles and permissions
# --------------------------------------------------------------------------
ROLES = {
    "admin": "Administrator",
    "head_teacher": "Head Teacher / Principal",
    "deputy_head_teacher": "Deputy Head Teacher",
    "exam_officer": "Examination officer",
    "teacher": "Teacher",
    "viewer": "Viewer (read-only)",
}

# Permissions the administrator can switch on/off per role.
# (The administrator always has everything, plus users, roles, settings,
# grading scale and restore, which are never delegated.)
PERMS = [
    ("students.view", "View the student list"),
    ("students.manage", "Add, edit, import and delete students"),
    ("classes.manage", "Manage classes and streams"),
    ("subjects.manage", "Manage subjects and teaching assignments"),
    ("exams.manage", "Create and manage examinations"),
    ("marks.enter", "Enter marks for their own assigned classes and subjects"),
    ("marks.enter_all", "Enter or edit marks for any class and subject"),
    ("marks.verify", "Verify and unlock marks"),
    ("results.view_all", "View results of all classes and subjects"),
    ("reports.print_all", "Generate report cards for all classes"),
    ("backup.create", "Create and download backups"),
    ("audit.view", "View the activity log"),
]

DEFAULT_ROLE_PERMS = {
    "head_teacher": ["students.view", "results.view_all", "reports.print_all", "marks.verify", "backup.create", "audit.view"],
    "deputy_head_teacher": ["students.view", "classes.manage", "subjects.manage", "exams.manage", "marks.enter", "marks.enter_all", "marks.verify", "results.view_all", "reports.print_all", "audit.view"],
    "exam_officer": [p for p, _ in PERMS],
    "teacher": ["marks.enter"],
    "viewer": ["students.view", "results.view_all", "reports.print_all"],
}

# --------------------------------------------------------------------------
# Grading presets
# --------------------------------------------------------------------------
GRADING_PRESETS = {
    "twelve": (
        "12-point scale (A to E)",
        [
            (80, "A", 12, "Excellent", "Excellent results. Keep up the outstanding work."),
            (75, "A-", 11, "Very good", "Very good performance. Aim even higher."),
            (70, "B+", 10, "Good", "Good performance. Keep working hard."),
            (65, "B", 9, "Good", "A good result. Push for further improvement."),
            (60, "B-", 8, "Fairly good", "A fairly good result. More effort will bring improvement."),
            (55, "C+", 7, "Fair", "A fair performance. Put in more effort."),
            (50, "C", 6, "Fair", "An average result. You can do better with more effort."),
            (45, "C-", 5, "Average", "Average performance. Work harder and seek help where needed."),
            (40, "D+", 4, "Below average", "Below average. Serious effort is needed."),
            (35, "D", 3, "Weak", "Weak performance. Work much harder."),
            (30, "D-", 2, "Weak", "Weak performance. Seek extra help and revise regularly."),
            (0, "E", 1, "Very weak", "Very poor performance. Urgent improvement is required."),
        ],
    ),
    "five": (
        "Simple scale (A to E)",
        [
            (80, "A", 5, "Excellent", "Excellent results. Keep up the outstanding work."),
            (65, "B", 4, "Good", "A good result. Keep working hard."),
            (50, "C", 3, "Fair", "A fair result. More effort will bring improvement."),
            (35, "D", 2, "Weak", "Weak performance. Work much harder."),
            (0, "E", 1, "Very weak", "Very poor performance. Urgent improvement is required."),
        ],
    ),
    "levels": (
        "Four performance levels (EE, ME, AE, BE)",
        [
            (75, "EE", 4, "Exceeding expectations", "Exceeding expectations. Keep up the excellent work."),
            (50, "ME", 3, "Meeting expectations", "Meeting expectations. Keep working hard."),
            (25, "AE", 2, "Approaching expectations", "Approaching expectations. More effort is needed."),
            (0, "BE", 1, "Below expectations", "Below expectations. Extra support and effort are needed."),
        ],
    ),
}

DEFAULT_SETTINGS = {
    "school_name": "Your School Name",
    "school_motto": "",
    "school_address": "",
    "school_phone": "",
    "school_email": "",
    "principal_name": "",
    "pass_mark": "40",
    "show_subject_position": "1",
    "require_verified": "0",
    "report_footer": "This report card is not valid without the school stamp.",
    "backup_copy_dir": "",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    full_name TEXT NOT NULL,
    tsc_no TEXT NOT NULL DEFAULT '',
    initials TEXT NOT NULL DEFAULT '',
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'teacher',
    phone TEXT NOT NULL DEFAULT '',
    active INTEGER NOT NULL DEFAULT 1,
    must_change_password INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL DEFAULT '');
CREATE TABLE IF NOT EXISTS role_permissions(
    role TEXT NOT NULL, perm TEXT NOT NULL, PRIMARY KEY(role, perm)
);
CREATE TABLE IF NOT EXISTS classes(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    stream TEXT NOT NULL DEFAULT '',
    class_teacher_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
    UNIQUE(name, stream)
);
CREATE TABLE IF NOT EXISTS subjects(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT NOT NULL UNIQUE COLLATE NOCASE,
    name TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS teaching(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    class_id INTEGER NOT NULL REFERENCES classes(id) ON DELETE CASCADE,
    subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
    UNIQUE(class_id, subject_id)
);
CREATE TABLE IF NOT EXISTS students(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    adm_no TEXT NOT NULL UNIQUE COLLATE NOCASE,
    uli TEXT NOT NULL DEFAULT '',
    assessment_no TEXT NOT NULL DEFAULT '',
    full_name TEXT NOT NULL,
    gender TEXT NOT NULL DEFAULT '',
    dob TEXT NOT NULL DEFAULT '',
    class_id INTEGER REFERENCES classes(id) ON DELETE SET NULL,
    guardian TEXT NOT NULL DEFAULT '',
    phone TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'active',
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS grade_bands(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    min_pct REAL NOT NULL,
    grade TEXT NOT NULL,
    points REAL NOT NULL DEFAULT 0,
    remark TEXT NOT NULL DEFAULT '',
    comment TEXT NOT NULL DEFAULT ''
);
CREATE TABLE IF NOT EXISTS exams(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    term TEXT NOT NULL DEFAULT '',
    year INTEGER,
    out_of REAL NOT NULL DEFAULT 100,
    status TEXT NOT NULL DEFAULT 'open',
    closing_date TEXT NOT NULL DEFAULT '',
    opening_date TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS exam_papers(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    exam_id INTEGER NOT NULL REFERENCES exams(id) ON DELETE CASCADE,
    class_id INTEGER NOT NULL REFERENCES classes(id) ON DELETE CASCADE,
    subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
    out_of REAL NOT NULL DEFAULT 100,
    UNIQUE(exam_id, class_id, subject_id)
);
CREATE TABLE IF NOT EXISTS marks(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    exam_id INTEGER NOT NULL REFERENCES exams(id) ON DELETE CASCADE,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    class_id INTEGER NOT NULL,
    subject_id INTEGER NOT NULL,
    score REAL,
    absent INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'entered',
    entered_by INTEGER,
    updated_at TEXT,
    UNIQUE(exam_id, student_id, subject_id)
);
CREATE INDEX IF NOT EXISTS idx_marks_paper ON marks(exam_id, class_id, subject_id);
CREATE TABLE IF NOT EXISTS report_remarks(
    exam_id INTEGER NOT NULL REFERENCES exams(id) ON DELETE CASCADE,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    class_teacher TEXT NOT NULL DEFAULT '',
    principal TEXT NOT NULL DEFAULT '',
    PRIMARY KEY(exam_id, student_id)
);
CREATE TABLE IF NOT EXISTS audit_log(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,
    user_id INTEGER,
    username TEXT NOT NULL DEFAULT '',
    action TEXT NOT NULL,
    detail TEXT NOT NULL DEFAULT ''
);
"""

CLASS_LABEL = ("(c.name || CASE WHEN COALESCE(c.stream,'')<>'' "
               "THEN ' ' || c.stream ELSE '' END)")


# --------------------------------------------------------------------------
# Database
# --------------------------------------------------------------------------
def connect():
    con = sqlite3.connect(DB_PATH, timeout=20)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con


def db():
    if "db" not in g:
        g.db = connect()
    return g.db


def close_db(_exc=None):
    con = g.pop("db", None)
    if con is not None:
        con.close()


def q(sql, args=()):
    return db().execute(sql, args).fetchall()


def q1(sql, args=()):
    return db().execute(sql, args).fetchone()


def ex(sql, args=()):
    return db().execute(sql, args)


def commit():
    db().commit()


def now():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def load_bands_into(con, preset_key):
    con.execute("DELETE FROM grade_bands")
    for m, g_, p, r, c in GRADING_PRESETS[preset_key][1]:
        con.execute("INSERT INTO grade_bands(min_pct,grade,points,remark,comment) "
                    "VALUES(?,?,?,?,?)", (m, g_, p, r, c))


def init_db():
    for d in (DATA_DIR, BACKUP_DIR, TMP_DIR):
        os.makedirs(d, exist_ok=True)
    con = connect()
    con.executescript(SCHEMA)
    if not con.execute("SELECT 1 FROM users").fetchone():
        con.execute(
            "INSERT INTO users(username,full_name,password_hash,role,must_change_password) "
            "VALUES('admin','System Administrator',?, 'admin', 1)",
            (generate_password_hash("admin123", method="pbkdf2:sha256"),))
    for k, v in DEFAULT_SETTINGS.items():
        con.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k, v))
    if not con.execute("SELECT 1 FROM grade_bands").fetchone():
        load_bands_into(con, "twelve")
    for role, perms in DEFAULT_ROLE_PERMS.items():
        for p in perms:
            con.execute("INSERT OR IGNORE INTO role_permissions(role,perm) VALUES(?,?)", (role, p))
    # Migrate students table if uli or assessment_no are missing
    cols = {c[1] for c in con.execute("PRAGMA table_info(students)").fetchall()}
    if cols and "uli" not in cols:
        con.execute("ALTER TABLE students ADD COLUMN uli TEXT NOT NULL DEFAULT ''")
    if cols and "assessment_no" not in cols:
        con.execute("ALTER TABLE students ADD COLUMN assessment_no TEXT NOT NULL DEFAULT ''")
    # Migrate users table if tsc_no is missing
    u_cols = {c[1] for c in con.execute("PRAGMA table_info(users)").fetchall()}
    if u_cols and "tsc_no" not in u_cols:
        con.execute("ALTER TABLE users ADD COLUMN tsc_no TEXT NOT NULL DEFAULT ''")
    con.commit()
    con.close()


def secret_key():
    if os.environ.get("SECRET_KEY"):
        return os.environ.get("SECRET_KEY").strip()
    path = os.path.join(DATA_DIR, "secret.key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


def hash_password(pw):
    return generate_password_hash(pw, method="pbkdf2:sha256")


def verify_password(h, pw):
    return check_password_hash(h, pw)


# --------------------------------------------------------------------------
# Settings, audit
# --------------------------------------------------------------------------
def settings():
    if "settings" not in g:
        d = dict(DEFAULT_SETTINGS)
        for r in q("SELECT key,value FROM settings"):
            d[r["key"]] = r["value"]
        g.settings = d
    return g.settings


def set_setting(key, value):
    ex("INSERT INTO settings(key,value) VALUES(?,?) "
       "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))
    g.pop("settings", None)


def audit(action, detail=""):
    u = g.get("user")
    ex("INSERT INTO audit_log(ts,user_id,username,action,detail) VALUES(?,?,?,?,?)",
       (now(), u["id"] if u else None, u["username"] if u else "", action, detail))


# --------------------------------------------------------------------------
# Sessions, CSRF and permissions
# --------------------------------------------------------------------------
def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


def check_csrf():
    tok = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token") or ""
    if not tok or not hmac.compare_digest(tok, session.get("csrf", "")):
        abort(400, "Your session expired or the form was invalid. Go back, refresh the page and try again.")


def load_user():
    g.user = None
    uid = session.get("uid")
    if uid:
        u = q1("SELECT * FROM users WHERE id=? AND active=1", (uid,))
        if u:
            g.user = u
        else:
            session.pop("uid", None)


def has_perm(perm):
    u = g.get("user")
    if not u:
        return False
    if u["role"] == "admin":
        return True
    if "perms" not in g:
        g.perms = {r["perm"] for r in q(
            "SELECT perm FROM role_permissions WHERE role=?", (u["role"],))}
    return perm in g.perms


def is_admin():
    u = g.get("user")
    return bool(u and u["role"] == "admin")


def perm_required(*perms):
    def deco(f):
        @functools.wraps(f)
        def wrapper(*a, **k):
            if not g.user:
                return redirect(url_for("auth.login", next=request.full_path.rstrip("?")))
            if not any(has_perm(p) for p in perms):
                abort(403)
            return f(*a, **k)
        return wrapper
    return deco


def login_required(f):
    @functools.wraps(f)
    def wrapper(*a, **k):
        if not g.user:
            return redirect(url_for("auth.login", next=request.full_path.rstrip("?")))
        return f(*a, **k)
    return wrapper


def admin_required(f):
    @functools.wraps(f)
    def wrapper(*a, **k):
        if not g.user:
            return redirect(url_for("auth.login", next=request.full_path.rstrip("?")))
        if not is_admin():
            abort(403)
        return f(*a, **k)
    return wrapper


# ---- data scope rules -----------------------------------------------------
def teacher_papers():
    """Set of (class_id, subject_id) assigned to the current user."""
    u = g.get("user")
    if not u:
        return set()
    return {(r["class_id"], r["subject_id"]) for r in q(
        "SELECT class_id, subject_id FROM teaching WHERE teacher_id=?", (u["id"],))}


def can_enter_marks(class_id, subject_id):
    u = g.get("user")
    if not u:
        return False
    if u["role"] == "admin":
        return True
    if u["role"] == "teacher":
        return (class_id, subject_id) in teacher_papers()
    if has_perm("marks.enter_all"):
        return True
    if has_perm("marks.enter"):
        return (class_id, subject_id) in teacher_papers()
    return False


def visible_class_ids():
    """None means all classes; otherwise a set of class ids the user may see."""
    if has_perm("results.view_all"):
        return None
    ids = {r["class_id"] for r in q("SELECT class_id FROM teaching WHERE teacher_id=?", (g.user["id"],))}
    ids |= {r["id"] for r in q("SELECT id FROM classes WHERE class_teacher_id=?", (g.user["id"],))}
    return ids


def class_teacher_ids():
    return {r["id"] for r in q("SELECT id FROM classes WHERE class_teacher_id=?", (g.user["id"],))}


def can_view_class(class_id):
    ids = visible_class_ids()
    return ids is None or class_id in ids


def can_print_reports(class_id):
    if has_perm("reports.print_all"):
        return True
    return class_id in class_teacher_ids()


def printable_class_ids():
    if has_perm("reports.print_all"):
        return None
    return class_teacher_ids()


def safe_next(url):
    if url and url.startswith("/") and not url.startswith("//"):
        return url
    return None


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def to_float(text):
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    t = str(text).strip().replace(",", ".")
    if t == "":
        return None
    return float(t)


def norm_cell(v):
    """Excel cell -> clean string (whole floats lose the .0)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def read_table_file(file_storage):
    """Read an uploaded .xlsx, .xlsm, or .csv file and return rows as list of lists of strings."""
    import csv
    import io
    import openpyxl

    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if filename.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        target_sheet = None
        for cand in ("Students", "Teachers", "Users"):
            if cand in wb.sheetnames:
                target_sheet = cand
                break
        ws = wb[target_sheet] if target_sheet else wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append([norm_cell(c) for c in r])
        return rows
    elif filename.endswith((".csv", ".tsv", ".txt")):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            return []
        try:
            sample = "\n".join(lines[:5])
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            reader = csv.reader(lines, dialect)
        except Exception:
            reader = csv.reader(lines)
        return [[str(c).strip() for c in r] for r in reader]
    else:
        raise ValueError("Unsupported file format. Please upload an Excel (.xlsx) or CSV (.csv) file.")


def flash_errors(errors, limit=8):
    for e in errors[:limit]:
        flash(e, "error")
    if len(errors) > limit:
        flash(f"...and {len(errors) - limit} more problems.", "error")


# --------------------------------------------------------------------------
# Backups
# --------------------------------------------------------------------------
def make_backup(kind="manual"):
    """Create a zip with a consistent copy of the database plus readable Excel copies."""
    import reports  # local import to avoid a circular import
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(BACKUP_DIR, f"exam_backup_{stamp}_{kind}.zip")
    snap = os.path.join(TMP_DIR, f"snapshot_{stamp}.db")
    os.makedirs(TMP_DIR, exist_ok=True)
    src = connect()
    dst = sqlite3.connect(snap)
    src.backup(dst)
    dst.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(snap, "exams.db")
        z.writestr("students.xlsx", reports.students_workbook(src).getvalue())
        z.writestr("marks.xlsx", reports.marks_workbook(src).getvalue())
        z.writestr("README.txt",
                   "School Examination System backup\n"
                   f"Created: {now()}\n\n"
                   "exams.db     - full database (use Backup > Restore to bring it back)\n"
                   "students.xlsx / marks.xlsx - readable copies you can open in Excel\n")
    src.close()
    os.remove(snap)
    copy_note = ""
    copy_dir = (settings().get("backup_copy_dir") or "").strip() if g else ""
    if copy_dir:
        try:
            os.makedirs(copy_dir, exist_ok=True)
            shutil.copy2(path, copy_dir)
        except OSError as e:
            copy_note = f"Could not copy to {copy_dir}: {e}"
    if kind == "auto":
        autos = sorted(f for f in os.listdir(BACKUP_DIR) if f.endswith("_auto.zip"))
        for old in autos[:-30]:
            try:
                os.remove(os.path.join(BACKUP_DIR, old))
            except OSError:
                pass
    return path, copy_note


def list_backups():
    if not os.path.isdir(BACKUP_DIR):
        return []
    out = []
    for f in sorted(os.listdir(BACKUP_DIR), reverse=True):
        if f.endswith(".zip"):
            p = os.path.join(BACKUP_DIR, f)
            out.append({"name": f, "size": os.path.getsize(p),
                        "time": datetime.datetime.fromtimestamp(os.path.getmtime(p)).strftime("%Y-%m-%d %H:%M")})
    return out


_last_auto_check = {"day": None}


def auto_backup_if_due():
    """Make one automatic backup per day, the first time the system is used that day."""
    today = datetime.date.today().isoformat()
    if _last_auto_check["day"] == today:
        return
    _last_auto_check["day"] = today
    stamp = today.replace("-", "")
    have = any(f.startswith(f"exam_backup_{stamp}") for f in os.listdir(BACKUP_DIR)) \
        if os.path.isdir(BACKUP_DIR) else False
    if not have and q1("SELECT 1 FROM students LIMIT 1"):
        try:
            make_backup("auto")
        except Exception:  # never block the user because a backup failed
            pass


def restore_backup(zip_path):
    """Replace the live database with the one inside a backup zip."""
    with zipfile.ZipFile(zip_path) as z:
        if "exams.db" not in z.namelist():
            raise ValueError("This file is not a backup from this system (exams.db is missing).")
        tmp = os.path.join(TMP_DIR, "restore_candidate.db")
        with open(tmp, "wb") as f:
            f.write(z.read("exams.db"))
    try:
        con = sqlite3.connect(tmp)
        ok = con.execute("PRAGMA integrity_check").fetchone()[0]
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        con.close()
        if ok != "ok" or not {"users", "students", "marks", "exams"} <= tables:
            raise ValueError("The database inside this backup is damaged or incomplete.")
        make_backup("before_restore")
        close_db()
        shutil.copyfile(tmp, DB_PATH)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


def fnum(v, dp=2):
    """Marks and totals: no trailing zeros (45, 45.5, 45.25)."""
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return v
    v = round(float(v), dp)
    return str(int(v)) if v == int(v) else f"{v:.{dp}f}".rstrip("0").rstrip(".")


def f1(v):
    """Percentages, means and points: always one decimal."""
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return v
    return f"{float(v):.1f}"
